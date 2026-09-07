"""data/history.json 전체 이력을 엑셀(xlsx)로 변환."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from storage import METRICS, METRIC_LABELS


def export_history_to_xlsx(data: dict[str, Any], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "금리이력"

    header = ["날짜"] + [METRIC_LABELS[m] for m in METRICS]
    ws.append(header)
    bold = Font(bold=True)
    fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for col_idx in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = bold
        cell.fill = fill

    for iso_date in sorted(data["days"].keys()):
        day = data["days"][iso_date]
        row = [iso_date]
        for m in METRICS:
            val = day["values"].get(m)
            row.append(val if val is not None else "")
        ws.append(row)

    for col_idx, _ in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


if __name__ == "__main__":
    import json
    import sys

    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("docs/data/history.json")
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("docs/data/history.xlsx")
    with src.open(encoding="utf-8") as f:
        history = json.load(f)
    export_history_to_xlsx(history, dst)
    print(f"saved {dst}")
